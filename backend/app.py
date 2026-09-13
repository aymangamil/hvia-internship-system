from __future__ import annotations

import io
import json
import os
import re
import secrets
import shutil
from datetime import datetime, timedelta
from pathlib import Path
from typing import Optional
from openpyxl import load_workbook
from urllib.parse import quote

from fastapi import FastAPI, Depends, HTTPException, Request, UploadFile, File, Form, Query
from fastapi.responses import FileResponse, HTMLResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from sqlalchemy import func, or_, and_
from sqlalchemy.orm import Session
from starlette.middleware.sessions import SessionMiddleware
from starlette.middleware.gzip import GZipMiddleware


from .database import init_db, get_db, SessionLocal
from .models import (
    User, Track, Enrollment, MentorAssignment, Task, TaskAttachment, TaskAssignment, TaskProgress,
    Submission, SubmissionFile, WorkSession, Meeting, MeetingAttendee, Announcement,
    LearningResource, ResourceProgress, Notification, AdminNote, Complaint, Warning,
    LoginEvent, ActivityLog, DemoProduct, ContactRequest, SystemSetting, UserPreference
)
from .security import verify_password, hash_password
from .seed import seed_database, add_attendees_for_meeting

ROOT = Path(__file__).resolve().parents[1]
FRONTEND = ROOT / "frontend"
UPLOADS = Path(os.getenv("UPLOADS_DIR", "/tmp/hvia_uploads" if os.getenv("VERCEL") else str(ROOT / "uploads")))
for p in [UPLOADS, UPLOADS/"task_attachments", UPLOADS/"submissions", UPLOADS/"profiles", UPLOADS/"resources"]:
    p.mkdir(parents=True, exist_ok=True)

app = FastAPI(title="HVIA Internship & Operations Management System", version="5.0.0")
app.add_middleware(GZipMiddleware, minimum_size=900)
app.add_middleware(SessionMiddleware, secret_key=os.getenv("SESSION_SECRET", "change-this-hvia-enterprise-session-secret-2026"), same_site="lax", https_only=os.getenv("COOKIE_HTTPS_ONLY", "false").lower() == "true")
app.mount("/assets", StaticFiles(directory=str(FRONTEND / "assets")), name="assets")


@app.on_event("startup")
def startup():
    init_db()
    db = SessionLocal()
    try:
        seed_database(db)
    finally:
        db.close()


# ---------- Input models ----------
class LoginIn(BaseModel):
    email: str
    password: str

class ProfileIn(BaseModel):
    phone: str = ""
    age: int | None = None
    linkedin_url: str = ""
    personal_email: str = ""
    university: str = ""
    faculty: str = ""
    graduation_year: int | None = None
    bio: str = ""
    track_id: int | None = None

class UserCreateIn(BaseModel):
    full_name: str
    email: str
    password: str
    personal_email: str | None = None

class AdminCreateIn(BaseModel):
    full_name: str
    email: str
    password: str

class UserUpdateIn(BaseModel):
    full_name: str | None = None
    email: str | None = None
    personal_email: str | None = None
    password: str | None = None
    phone: str | None = None
    age: int | None = None
    linkedin_url: str | None = None
    university: str | None = None
    faculty: str | None = None
    graduation_year: int | None = None
    status: str | None = None
    active: bool | None = None

class AdminUpdateIn(BaseModel):
    full_name: str | None = None
    email: str | None = None
    password: str | None = None
    status: str | None = None
    active: bool | None = None

class TaskIn(BaseModel):
    title: str
    description: str = ""
    instructions: str = ""
    target_type: str = "all"
    target_track_id: int | None = None
    target_user_id: int | None = None
    # Optional multi-select. When present, TaskAssignment rows become the source of truth.
    target_user_ids: list[int] = []
    priority: str = "Medium"
    deadline: datetime | None = None
    max_score: int = 100
    grace_minutes: int = 0
    allow_resubmission: bool = True

class MeetingIn(BaseModel):
    title: str
    description: str = ""
    starts_at: datetime
    duration_minutes: int = 60
    meeting_link: str = ""
    target_type: str = "interns"
    target_track_id: int | None = None
    target_user_id: int | None = None

class AnnouncementIn(BaseModel):
    title: str
    content: str
    target_type: str = "all"
    target_track_id: int | None = None
    target_role: str | None = None
    pinned: bool = False
    publish_at: datetime | None = None

class ReviewIn(BaseModel):
    score: float | None = None
    feedback: str = ""
    status: str = "Completed"

class ComplaintIn(BaseModel):
    subject: str
    message: str
    category: str = "General"
    priority: str = "Normal"

class ComplaintReplyIn(BaseModel):
    status: str = "Resolved"
    reply: str = ""

class AdminNoteIn(BaseModel):
    admin_id: int
    title: str
    content: str
    priority: str = "Normal"

class WarningIn(BaseModel):
    user_id: int
    warning_type: str
    severity: str = "Minor"
    note: str

class ContactIn(BaseModel):
    name: str
    phone: str
    email: str | None = None
    company: str | None = None
    interest: str | None = None
    message: str | None = None

class SettingIn(BaseModel):
    value: str

class TrackChangeIn(BaseModel):
    track_id: int

class MentorAssignIn(BaseModel):
    intern_ids: list[int]

class AttendanceIn(BaseModel):
    status: str
    note: str = ""

class ResourceProgressIn(BaseModel):
    watched_seconds: int = 0
    progress_percent: float = 0
    completed: bool = False


# ---------- Helpers ----------
def active_enrollment(db: Session, user_id: int):
    return db.query(Enrollment).filter(Enrollment.user_id == user_id, Enrollment.active == True).order_by(Enrollment.id.desc()).first()


def serialize_user(db: Session, u: User):
    e = active_enrollment(db, u.id)
    track = None
    if e and e.track_id:
        t = db.get(Track, e.track_id)
        if t:
            track = {"id": t.id, "name": t.name, "program_type": t.program_type}
    return {
        "id": u.id, "internal_id": u.internal_id, "full_name": u.full_name, "email": u.email,
        "personal_email": u.personal_email, "role": u.role, "status": u.status, "active": u.active,
        "profile_completed": u.profile_completed, "phone": u.phone, "age": u.age,
        "linkedin_url": u.linkedin_url, "university": u.university, "faculty": u.faculty,
        "graduation_year": u.graduation_year, "bio": u.bio, "profile_picture": u.profile_picture,
        "last_login_at": u.last_login_at, "last_seen_at": u.last_seen_at, "created_at": u.created_at,
        "track": track, "enrollment_status": e.status if e else None,
    }


def current_user(request: Request, db: Session = Depends(get_db)) -> User:
    uid = request.session.get("user_id")
    if not uid:
        raise HTTPException(401, "Not authenticated")
    u = db.get(User, uid)
    if not u or not u.active or u.deleted_at is not None:
        request.session.clear()
        raise HTTPException(401, "Account unavailable")
    u.last_seen_at = datetime.utcnow()
    db.commit()
    return u


def require_roles(*roles):
    def dep(u: User = Depends(current_user)):
        if u.role not in roles:
            raise HTTPException(403, "Not allowed")
        return u
    return dep


def ip_of(request: Request):
    return request.client.host if request.client else None


def log(db: Session, actor: Optional[User], action: str, entity_type=None, entity_id=None, details=None, request: Request | None = None):
    db.add(ActivityLog(
        actor_user_id=actor.id if actor else None, action=action, entity_type=entity_type,
        entity_id=entity_id, details=json.dumps(details, ensure_ascii=False, default=str) if isinstance(details, (dict, list)) else details,
        ip_address=ip_of(request) if request else None,
    ))


def notify(db: Session, user_id: int, title: str, message: str, kind="general", related_type=None, related_id=None):
    db.add(Notification(user_id=user_id, title=title, message=message, notification_type=kind, related_type=related_type, related_id=related_id))


def notify_many(db: Session, user_ids: list[int], title: str, message: str, kind="general", related_type=None, related_id=None):
    for uid in set(user_ids):
        notify(db, uid, title, message, kind, related_type, related_id)


def safe_name(name: str):
    base = Path(name or "file").name
    base = re.sub(r"[^A-Za-z0-9._()\- ]+", "_", base).strip() or "file"
    return base[:220]


def save_upload(upload: UploadFile, folder: Path):
    original = safe_name(upload.filename or "file")
    token = secrets.token_hex(8)
    stored = f"{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{token}_{original}"
    path = folder / stored
    with path.open("wb") as f:
        shutil.copyfileobj(upload.file, f)
    return original, stored, path.stat().st_size


def mentor_intern_ids(db: Session, admin_id: int):
    return [x.intern_id for x in db.query(MentorAssignment).filter(MentorAssignment.admin_id == admin_id, MentorAssignment.active == True).all()]


def can_access_intern(db: Session, actor: User, intern_id: int):
    if actor.role == "super_admin": return True
    if actor.role == "admin": return intern_id in mentor_intern_ids(db, actor.id)
    return actor.id == intern_id


def target_users_for_task(db: Session, task: Task, actor: User):
    q = db.query(User).filter(User.role == "intern", User.active == True, User.deleted_at.is_(None))
    if actor.role == "admin":
        scope = mentor_intern_ids(db, actor.id)
        q = q.filter(User.id.in_(scope or [-1]))
    if task.target_type == "specific" and task.target_user_id:
        q = q.filter(User.id == task.target_user_id)
    elif task.target_type == "track" and task.target_track_id:
        ids = [e.user_id for e in db.query(Enrollment).filter(Enrollment.track_id == task.target_track_id, Enrollment.active == True).all()]
        q = q.filter(User.id.in_(ids or [-1]))
    return q.all()


def task_for_user(db: Session, task_id: int, user_id: int):
    return db.query(TaskAssignment).filter(TaskAssignment.task_id == task_id, TaskAssignment.user_id == user_id).first() is not None


def assign_existing_tasks_to_intern(db: Session, intern_id: int, track_id: int | None):
    q = db.query(Task).filter(Task.archived_at.is_(None))
    tasks = q.filter(or_(Task.target_type == "all", and_(Task.target_type == "track", Task.target_track_id == track_id), and_(Task.target_type == "specific", Task.target_user_id == intern_id))).all()
    for task in tasks:
        if not db.query(TaskAssignment).filter(TaskAssignment.task_id == task.id, TaskAssignment.user_id == intern_id).first():
            db.add(TaskAssignment(task_id=task.id, user_id=intern_id))
            db.add(TaskProgress(task_id=task.id, user_id=intern_id, status="To Do"))
            notify(db, intern_id, "Task Assigned", task.title, "task", "task", task.id)


def latest_submission(db: Session, user_id: int, task_id: int):
    return db.query(Submission).filter(Submission.user_id == user_id, Submission.task_id == task_id).order_by(Submission.version.desc()).first()


def submission_files(db: Session, submission_id: int):
    return db.query(SubmissionFile).filter(SubmissionFile.submission_id == submission_id).all()


def serialize_submission(db: Session, s: Submission):
    task = db.get(Task, s.task_id)
    user = db.get(User, s.user_id)
    reviewer = db.get(User, s.reviewed_by) if s.reviewed_by else None
    return {
        "id": s.id, "task_id": s.task_id, "task": {"id": task.id, "title": task.title, "max_score": task.max_score} if task else None,
        "user_id": s.user_id, "student": user.full_name if user else "—", "student_email": user.email if user else "—",
        "version": s.version, "notes": s.notes, "github_url": s.github_url, "drive_url": s.drive_url,
        "status": s.status, "submitted_at": s.submitted_at, "is_late": s.is_late, "score": s.score,
        "feedback": s.feedback, "reviewed_by": reviewer.full_name if reviewer else None, "reviewed_at": s.reviewed_at,
        "files": [{"id": f.id, "name": f.original_name, "size": f.file_size, "content_type": f.content_type, "download_url": f"/api/submission-files/{f.id}/download"} for f in submission_files(db, s.id)],
    }


def task_stats_for_intern(db: Session, user_id: int):
    assignments = db.query(TaskAssignment).filter(TaskAssignment.user_id == user_id).all()
    ids = [a.task_id for a in assignments]
    total = len(ids)
    completed = db.query(TaskProgress).filter(TaskProgress.user_id == user_id, TaskProgress.task_id.in_(ids or [-1]), TaskProgress.status == "Completed").count()
    submitted = db.query(TaskProgress).filter(TaskProgress.user_id == user_id, TaskProgress.task_id.in_(ids or [-1]), TaskProgress.status.in_(["Submitted", "Under Review", "Completed", "Changes Requested"])).count()
    late = db.query(Submission).filter(Submission.user_id == user_id, Submission.is_late == True).count()
    scores = [x[0] for x in db.query(Submission.score).filter(Submission.user_id == user_id, Submission.score.isnot(None)).all()]
    avg = round(sum(scores)/len(scores), 1) if scores else 0
    return {"total": total, "completed": completed, "submitted": submitted, "late": late, "avg_score": avg, "completion_percent": round(completed/total*100, 1) if total else 0}


def hours_for_user(db: Session, user_id: int, start: datetime | None = None, end: datetime | None = None):
    q = db.query(func.coalesce(func.sum(WorkSession.minutes), 0)).filter(WorkSession.user_id == user_id)
    if start: q = q.filter(WorkSession.started_at >= start)
    if end: q = q.filter(WorkSession.started_at <= end)
    minutes = int(q.scalar() or 0)
    return round(minutes / 60, 2)


def attendance_for_user(db: Session, user_id: int):
    rows = db.query(MeetingAttendee).filter(MeetingAttendee.user_id == user_id).all()
    if not rows: return {"total": 0, "present": 0, "percent": 0}
    present = sum(1 for x in rows if x.attendance_status in ("present", "late"))
    return {"total": len(rows), "present": present, "percent": round(present/len(rows)*100, 1)}


def intern_report_data(db: Session, intern_id: int, start: datetime | None = None, end: datetime | None = None):
    u = db.get(User, intern_id)
    if not u or u.role != "intern": raise HTTPException(404, "Intern not found")
    stats = task_stats_for_intern(db, intern_id)
    att = attendance_for_user(db, intern_id)
    warnings = db.query(Warning).filter(Warning.user_id == intern_id).count()
    subsq = db.query(Submission).filter(Submission.user_id == intern_id)
    if start: subsq = subsq.filter(Submission.submitted_at >= start)
    if end: subsq = subsq.filter(Submission.submitted_at <= end)
    return {"user": serialize_user(db, u), "tasks": stats, "attendance": att, "hours": hours_for_user(db, intern_id, start, end), "warnings": warnings, "submissions_period": subsq.count()}


def parse_dt_date(value: str | None, end=False):
    if not value: return None
    try:
        d = datetime.fromisoformat(value)
        if len(value) <= 10:
            d = d.replace(hour=23, minute=59, second=59) if end else d.replace(hour=0, minute=0, second=0)
        return d
    except Exception:
        raise HTTPException(400, "Invalid date")


def get_setting(db: Session, key: str, default=""):
    s = db.query(SystemSetting).filter(SystemSetting.key == key).first()
    return s.value if s else default


# ---------- Public ----------
@app.get("/", response_class=HTMLResponse)
def index():
    return (FRONTEND / "index.html").read_text(encoding="utf-8")

@app.get("/api/public/demos")
def public_demos(db: Session = Depends(get_db)):
    return [{"id": x.id, "name": x.name, "category": x.category, "description": x.description, "demo_url": x.demo_url, "image_url": x.image_url} for x in db.query(DemoProduct).filter(DemoProduct.active == True).all()]

@app.post("/api/public/contact")
def public_contact(payload: ContactIn, db: Session = Depends(get_db)):
    r = ContactRequest(**payload.model_dump())
    db.add(r); db.commit(); db.refresh(r)
    number = get_setting(db, "whatsapp_number", "201222970033")
    msg = f"HVIA Demo Request\nName: {payload.name}\nPhone: {payload.phone}\nCompany: {payload.company or '-'}\nInterest: {payload.interest or '-'}\nMessage: {payload.message or '-'}"
    return {"ok": True, "request_id": r.id, "whatsapp_url": f"https://wa.me/{number}?text={quote(msg)}"}


# ---------- Auth ----------
@app.post("/api/auth/login")
def login_api(payload: LoginIn, request: Request, db: Session = Depends(get_db)):
    now = datetime.utcnow()
    recent_failures = db.query(LoginEvent).filter(LoginEvent.email_attempt == payload.email.lower().strip(), LoginEvent.success == False, LoginEvent.created_at >= now - timedelta(minutes=15)).count()
    if recent_failures >= 5:
        raise HTTPException(429, "Too many failed attempts. Try again later.")
    u = db.query(User).filter(func.lower(User.email) == payload.email.lower().strip()).first()
    ok = bool(u and u.active and u.deleted_at is None and verify_password(payload.password, u.password_hash))
    db.add(LoginEvent(user_id=u.id if u else None, email_attempt=payload.email.lower().strip(), success=ok, ip_address=ip_of(request), user_agent=request.headers.get("user-agent", "")[:500]))
    if not ok:
        db.commit(); raise HTTPException(401, "Invalid email or password")
    request.session["user_id"] = u.id
    u.last_login_at = now; u.last_seen_at = now
    log(db, u, "LOGIN", "user", u.id, {"role": u.role}, request)
    db.commit()
    return serialize_user(db, u)

@app.post("/api/auth/logout")
def logout_api(request: Request, db: Session = Depends(get_db), u: User = Depends(current_user)):
    log(db, u, "LOGOUT", "user", u.id, request=request); db.commit(); request.session.clear(); return {"ok": True}

@app.get("/api/auth/me")
def me(db: Session = Depends(get_db), u: User = Depends(current_user)):
    return serialize_user(db, u)

@app.get("/api/auth/login-history")
def my_login_history(db: Session = Depends(get_db), u: User = Depends(current_user)):
    return [{"success": x.success, "ip": x.ip_address, "user_agent": x.user_agent, "created_at": x.created_at} for x in db.query(LoginEvent).filter(LoginEvent.user_id == u.id).order_by(LoginEvent.id.desc()).limit(30).all()]


# ---------- Profiles & tracks ----------
@app.get("/api/tracks")
def tracks(db: Session = Depends(get_db), u: User = Depends(current_user)):
    return [{"id": t.id, "name": t.name, "program_type": t.program_type, "description": t.description, "active": t.active, "capacity": t.capacity} for t in db.query(Track).filter(Track.active == True).all()]

@app.post("/api/profile/complete")
def complete_profile(payload: ProfileIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("intern"))):
    for k in ["phone", "age", "linkedin_url", "personal_email", "university", "faculty", "graduation_year", "bio"]:
        setattr(u, k, getattr(payload, k))
    u.profile_completed = True; u.status = "active"
    if payload.track_id:
        t = db.get(Track, payload.track_id)
        if not t or not t.active: raise HTTPException(400, "Invalid track")
        old = active_enrollment(db, u.id)
        if old: old.active = False
        db.add(Enrollment(user_id=u.id, track_id=t.id, status="active", active=True, enrolled_at=datetime.utcnow()))
        db.flush()
        assign_existing_tasks_to_intern(db, u.id, t.id)
    log(db, u, "PROFILE_COMPLETED", "user", u.id, {"track_id": payload.track_id}, request)
    db.commit(); return serialize_user(db, u)

@app.post("/api/profile/photo")
async def upload_profile_photo(upload: UploadFile = File(...), request: Request = None, db: Session = Depends(get_db), u: User = Depends(current_user)):
    if not (upload.content_type or "").startswith("image/"):
        raise HTTPException(400, "Profile picture must be an image")
    suffix = Path(upload.filename or "photo.jpg").suffix.lower()[:10] or ".jpg"
    stored = f"profile_{u.id}{suffix}"
    path = UPLOADS / "profiles" / stored
    with path.open("wb") as f:
        shutil.copyfileobj(upload.file, f)
    u.profile_picture = f"/api/profile/photo/{u.id}"
    log(db, u, "PROFILE_PHOTO_UPDATED", "user", u.id, {"file": safe_name(upload.filename or stored), "size": path.stat().st_size}, request)
    db.commit(); return {"ok": True, "url": u.profile_picture}

@app.get("/api/profile/photo/{user_id}")
def profile_photo(user_id: int, db: Session = Depends(get_db)):
    u = db.get(User, user_id)
    if not u or not u.profile_picture: raise HTTPException(404)
    files = sorted((UPLOADS/"profiles").glob(f"profile_{user_id}.*"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not files: raise HTTPException(404)
    return FileResponse(files[0])


# ---------- People / admins ----------
@app.get("/api/admins")
def list_admins(db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin"))):
    out = []
    for a in db.query(User).filter(User.role == "admin", User.deleted_at.is_(None)).all():
        assigned = mentor_intern_ids(db, a.id)
        reviews = db.query(Submission).filter(Submission.reviewed_by == a.id).count()
        tasks = db.query(Task).filter(Task.created_by == a.id).count()
        meetings = db.query(Meeting).filter(Meeting.created_by == a.id).count()
        out.append({**serialize_user(db, a), "assigned_interns": len(assigned), "assigned_intern_ids": assigned, "reviews": reviews, "tasks_created": tasks, "meetings_created": meetings})
    return out

@app.post("/api/admins")
def add_admin(payload: AdminCreateIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin"))):
    if db.query(User).filter(func.lower(User.email) == payload.email.lower()).first(): raise HTTPException(400, "Email already exists")
    idx = db.query(User).filter(User.role == "admin").count() + 1
    a = User(internal_id=f"HVIA-ADM-2026-{idx+2:03d}", full_name=payload.full_name, email=payload.email.lower(), password_hash=hash_password(payload.password), role="admin", active=True, status="active", profile_completed=True)
    db.add(a); db.flush(); log(db, u, "ADMIN_CREATED", "user", a.id, {"email": a.email}, request); db.commit(); return serialize_user(db, a)

@app.patch("/api/admins/{admin_id}")
def update_admin(admin_id: int, payload: AdminUpdateIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin"))):
    a = db.get(User, admin_id)
    if not a or a.role != "admin" or a.deleted_at is not None: raise HTTPException(404, "Admin not found")
    if payload.email and payload.email.lower() != a.email.lower():
        if db.query(User).filter(func.lower(User.email) == payload.email.lower(), User.id != a.id).first(): raise HTTPException(400, "Email already exists")
        a.email = payload.email.lower().strip()
    if payload.full_name is not None: a.full_name = payload.full_name.strip()
    if payload.password: a.password_hash = hash_password(payload.password)
    if payload.status is not None:
        a.status = payload.status
        if payload.active is None:
            a.active = payload.status not in {"terminated", "inactive", "archived"}
    if payload.active is not None: a.active = payload.active
    log(db, u, "ADMIN_UPDATED", "user", a.id, {"email": a.email, "status": a.status, "active": a.active}, request)
    db.commit(); return serialize_user(db, a)

@app.get("/api/admin-options")
def admin_options(db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin"))):
    return [{"id":a.id,"full_name":a.full_name,"email":a.email,"status":a.status} for a in db.query(User).filter(User.role=="admin",User.active==True,User.deleted_at.is_(None)).order_by(User.full_name).all()]

@app.delete("/api/admins/{admin_id}")
def delete_admin(admin_id: int, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin"))):
    a = db.get(User, admin_id)
    if not a or a.role != "admin": raise HTTPException(404)
    affected=[x.intern_id for x in db.query(MentorAssignment).filter(MentorAssignment.admin_id==a.id,MentorAssignment.active==True).all()]
    for row in db.query(MentorAssignment).filter(MentorAssignment.admin_id==a.id,MentorAssignment.active==True).all(): row.active=False
    a.active = False; a.status = "terminated"; a.deleted_at = datetime.utcnow()
    notify_many(db,affected,"Monitor Assignment Updated","Your previous monitor account is no longer active. The Super Admin will assign a new monitor.","mentor")
    log(db, u, "ADMIN_ARCHIVED", "user", a.id, {"email": a.email,"affected_interns":affected}, request); db.commit(); return {"ok": True}

@app.post("/api/admins/{admin_id}/assign-interns")
def assign_admin_interns(admin_id: int, payload: MentorAssignIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin"))):
    a = db.get(User, admin_id)
    if not a or a.role != "admin": raise HTTPException(404)
    # Replace this admin's scope, and keep one active monitor per intern.
    for old in db.query(MentorAssignment).filter(MentorAssignment.admin_id == admin_id, MentorAssignment.active == True).all():
        old.active = False
    chosen = set(payload.intern_ids)
    valid_interns = db.query(User).filter(User.id.in_(chosen or [-1]), User.role == "intern", User.active == True, User.deleted_at.is_(None)).all() if chosen else []
    if {x.id for x in valid_interns} != chosen:
        raise HTTPException(400, "One or more intern IDs are invalid or inactive")
    if chosen:
        for old in db.query(MentorAssignment).filter(MentorAssignment.intern_id.in_(chosen), MentorAssignment.active == True).all():
            old.active = False
    for intern in valid_interns:
        db.add(MentorAssignment(admin_id=admin_id, intern_id=intern.id, assigned_by=u.id, active=True))
        notify(db, intern.id, "Monitor Assigned", f"{a.full_name} is now your HVIA monitor / mentor.", "mentor")
    log(db, u, "MENTOR_SCOPE_UPDATED", "user", admin_id, {"intern_ids": payload.intern_ids}, request); db.commit(); return {"ok": True}

@app.post("/api/admin-notes")
def create_admin_note(payload: AdminNoteIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin"))):
    a = db.get(User, payload.admin_id)
    if not a or a.role != "admin": raise HTTPException(404)
    n = AdminNote(admin_id=a.id, created_by=u.id, title=payload.title, content=payload.content, priority=payload.priority)
    db.add(n); db.flush(); notify(db, a.id, payload.title, payload.content, "admin_note", "admin_note", n.id)
    log(db, u, "ADMIN_NOTE_SENT", "admin_note", n.id, {"admin_id": a.id}, request); db.commit(); return {"ok": True}

@app.get("/api/admin-notes")
def admin_notes(db: Session = Depends(get_db), u: User = Depends(require_roles("admin", "super_admin"))):
    q = db.query(AdminNote)
    if u.role == "admin": q = q.filter(AdminNote.admin_id == u.id)
    rows = q.order_by(AdminNote.id.desc()).all()
    return [{"id": x.id, "admin_id": x.admin_id, "title": x.title, "content": x.content, "priority": x.priority, "acknowledged": x.acknowledged, "created_at": x.created_at} for x in rows]

@app.patch("/api/admin-notes/{note_id}/ack")
def ack_admin_note(note_id: int, db: Session = Depends(get_db), u: User = Depends(require_roles("admin"))):
    n = db.get(AdminNote, note_id)
    if not n or n.admin_id != u.id: raise HTTPException(404)
    n.acknowledged = True; db.commit(); return {"ok": True}

@app.get("/api/interns")
def list_interns(track_id: int | None = Query(None), db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin", "admin"))):
    q = db.query(User).filter(User.role == "intern", User.deleted_at.is_(None))
    if u.role == "admin":
        q = q.filter(User.id.in_(mentor_intern_ids(db, u.id) or [-1]))
    if track_id:
        scoped_ids = [e.user_id for e in db.query(Enrollment).filter(Enrollment.track_id == track_id, Enrollment.active == True).all()]
        q = q.filter(User.id.in_(scoped_ids or [-1]))
    out = []
    for x in q.order_by(User.id).all():
        ma = db.query(MentorAssignment).filter(MentorAssignment.intern_id == x.id, MentorAssignment.active == True).order_by(MentorAssignment.id.desc()).first()
        mentor = db.get(User, ma.admin_id) if ma else None
        out.append({**serialize_user(db, x), "mentor": ({"id": mentor.id, "name": mentor.full_name, "email": mentor.email} if mentor else None), "report": intern_report_data(db, x.id)})
    return out

@app.get("/api/interns/{intern_id}/summary")
def intern_summary(intern_id: int, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin", "admin"))):
    if not can_access_intern(db, u, intern_id):
        raise HTTPException(403, "Intern outside your assigned scope")
    return intern_report_data(db, intern_id)


@app.get("/api/intern-options")
def intern_options(track_id: int | None = Query(None), db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin", "admin"))):
    """Fast lightweight list for assignment selectors.

    A blank track returns all interns in the current user's allowed scope.
    Selecting a track returns only active enrollments in that track.
    """
    q = db.query(User).filter(User.role == "intern", User.active == True, User.deleted_at.is_(None))
    if u.role == "admin":
        q = q.filter(User.id.in_(mentor_intern_ids(db, u.id) or [-1]))
    if track_id:
        scoped_ids = [e.user_id for e in db.query(Enrollment).filter(Enrollment.track_id == track_id, Enrollment.active == True).all()]
        q = q.filter(User.id.in_(scoped_ids or [-1]))
    rows = q.order_by(User.full_name.asc()).all()
    out = []
    for x in rows:
        e = active_enrollment(db, x.id)
        t = db.get(Track, e.track_id) if e and e.track_id else None
        out.append({
            "id": x.id, "internal_id": x.internal_id, "full_name": x.full_name, "email": x.email,
            "track_id": t.id if t else None, "track_name": t.name if t else "Pending Track"
        })
    return out

@app.post("/api/interns")
def add_intern(payload: UserCreateIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin"))):
    if db.query(User).filter(func.lower(User.email) == payload.email.lower()).first(): raise HTTPException(400, "Email already exists")
    idx = db.query(User).filter(User.role == "intern").count() + 1
    intern = User(internal_id=f"HVIA-INT-2026-{idx:04d}", full_name=payload.full_name, email=payload.email.lower(), personal_email=payload.personal_email, password_hash=hash_password(payload.password), role="intern", active=True, status="onboarding", profile_completed=False)
    db.add(intern); db.flush(); db.add(Enrollment(user_id=intern.id, track_id=None, status="pending", active=True))
    log(db, u, "INTERN_CREATED", "user", intern.id, {"email": intern.email}, request); db.commit(); return serialize_user(db, intern)

@app.patch("/api/interns/{intern_id}")
def update_intern(intern_id: int, payload: UserUpdateIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin"))):
    intern = db.get(User, intern_id)
    if not intern or intern.role != "intern" or intern.deleted_at is not None: raise HTTPException(404, "Intern not found")
    if payload.email and payload.email.lower() != intern.email.lower():
        if db.query(User).filter(func.lower(User.email)==payload.email.lower(),User.id!=intern.id).first(): raise HTTPException(400,"Email already exists")
        intern.email=payload.email.lower().strip()
    for field in ["full_name","personal_email","phone","age","linkedin_url","university","faculty","graduation_year","status","active"]:
        value=getattr(payload,field)
        if value is not None: setattr(intern,field,value.strip() if isinstance(value,str) else value)
    if payload.status is not None and payload.active is None:
        intern.active = payload.status not in {"terminated", "dropped"}
    if payload.password: intern.password_hash=hash_password(payload.password)
    log(db,u,"INTERN_UPDATED","user",intern.id,{"email":intern.email,"status":intern.status},request);db.commit();return serialize_user(db,intern)

@app.get("/api/interns/import-template")
def download_import_template(u: User = Depends(require_roles("super_admin"))):
    path=ROOT/"data"/"HVIA_BULK_INTERN_IMPORT_TEMPLATE.xlsx"
    if not path.exists(): raise HTTPException(404,"Template not found")
    return FileResponse(path,filename=path.name,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.post("/api/interns/import-xlsx")
async def import_interns_xlsx(upload: UploadFile = File(...), request: Request = None, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin"))):
    if not (upload.filename or "").lower().endswith(".xlsx"): raise HTTPException(400,"Upload an .xlsx file")
    raw=await upload.read()
    if len(raw)>10*1024*1024: raise HTTPException(400,"Excel file is too large (max 10 MB)")
    try:
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True); ws=wb.active
    except Exception as ex: raise HTTPException(400,f"Invalid Excel file: {ex}")
    headers=[str(x.value or '').strip().lower() for x in ws[1]]
    aliases={"username":"username","email":"username","company_email":"username","password":"password","full_name":"full_name","name":"full_name","personal_email":"personal_email","track":"track","phone":"phone","age":"age","linkedin":"linkedin_url","linkedin_url":"linkedin_url","university":"university","faculty":"faculty","graduation_year":"graduation_year","status":"status","monitor_email":"monitor_email"}
    mapped=[aliases.get(h,h) for h in headers]
    missing=[h for h in ("username","password","full_name") if h not in mapped]
    if missing: raise HTTPException(400,"Missing required columns: "+", ".join(missing))
    rows=[];errors=[];seen=set()
    tracks_by_name={t.name.strip().lower():t for t in db.query(Track).filter(Track.active==True).all()}
    admins_by_email={a.email.lower():a for a in db.query(User).filter(User.role=="admin",User.active==True,User.deleted_at.is_(None)).all()}
    for ridx,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
        if not any(v not in (None,'') for v in row): continue
        data={mapped[i]:row[i] if i<len(row) else None for i in range(len(mapped))}
        email=str(data.get("username") or '').strip().lower(); password=str(data.get("password") or '').strip(); name=str(data.get("full_name") or '').strip()
        if not email or not password or not name: errors.append(f"Row {ridx}: username, password and full_name are required"); continue
        if email in seen: errors.append(f"Row {ridx}: duplicate username in file: {email}"); continue
        seen.add(email)
        if db.query(User).filter(func.lower(User.email)==email).first(): errors.append(f"Row {ridx}: user already exists: {email}"); continue
        track=None
        if data.get("track") not in (None,''):
            track=tracks_by_name.get(str(data.get("track")).strip().lower())
            if not track: errors.append(f"Row {ridx}: unknown track: {data.get('track')}"); continue
        monitor=None
        if data.get("monitor_email") not in (None,''):
            monitor=admins_by_email.get(str(data.get("monitor_email")).strip().lower())
            if not monitor: errors.append(f"Row {ridx}: unknown monitor_email: {data.get('monitor_email')}"); continue
        try:
            if data.get("age") not in (None, ''): data["age"] = int(data.get("age"))
            if data.get("graduation_year") not in (None, ''): data["graduation_year"] = int(data.get("graduation_year"))
        except (TypeError, ValueError):
            errors.append(f"Row {ridx}: age and graduation_year must be whole numbers"); continue
        status=str(data.get("status") or ("active" if track else "onboarding")).strip().lower()
        if status not in {"onboarding","active","on_hold","completed","dropped","terminated"}:
            errors.append(f"Row {ridx}: invalid status: {status}"); continue
        data["status"] = status
        rows.append((ridx,data,email,password,name,track,monitor))
    if errors: raise HTTPException(400,"Excel validation failed\n"+"\n".join(errors[:30]))
    created=[]
    for _,data,email,password,name,track,monitor in rows:
        idx=db.query(User).filter(User.role=="intern").count()+1
        internal=f"HVIA-INT-2026-{idx:04d}"
        while db.query(User).filter(User.internal_id==internal).first(): idx+=1;internal=f"HVIA-INT-2026-{idx:04d}"
        status=data["status"]
        person=User(internal_id=internal,full_name=name,email=email,personal_email=(str(data.get("personal_email")).strip() if data.get("personal_email") else None),password_hash=hash_password(password),role="intern",active=status not in {"terminated","dropped"},status=status,profile_completed=False,phone=(str(data.get("phone")).strip() if data.get("phone") else None),age=int(data.get("age")) if data.get("age") not in (None,'') else None,linkedin_url=(str(data.get("linkedin_url")).strip() if data.get("linkedin_url") else None),university=(str(data.get("university")).strip() if data.get("university") else None),faculty=(str(data.get("faculty")).strip() if data.get("faculty") else None),graduation_year=int(data.get("graduation_year")) if data.get("graduation_year") not in (None,'') else None)
        db.add(person);db.flush();db.add(Enrollment(user_id=person.id,track_id=track.id if track else None,status="active" if track else "pending",active=True,enrolled_at=datetime.utcnow() if track else None))
        if monitor: db.add(MentorAssignment(admin_id=monitor.id,intern_id=person.id,assigned_by=u.id,active=True))
        created.append(person)
    log(db,u,"INTERNS_BULK_IMPORTED","user",None,{"count":len(created),"file":safe_name(upload.filename or 'import.xlsx')},request);db.commit()
    return {"ok":True,"imported":len(created),"users":[{"id":x.id,"name":x.full_name,"email":x.email,"internal_id":x.internal_id} for x in created]}

@app.patch("/api/interns/{intern_id}/track")
def change_track(intern_id: int, payload: TrackChangeIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin"))):
    intern = db.get(User, intern_id); track = db.get(Track, payload.track_id)
    if not intern or intern.role != "intern" or not track: raise HTTPException(404)
    old = active_enrollment(db, intern_id)
    if old: old.active = False
    db.add(Enrollment(user_id=intern_id, track_id=track.id, status="active", active=True, enrolled_at=datetime.utcnow()))
    db.flush()
    assign_existing_tasks_to_intern(db, intern_id, track.id)
    notify(db, intern_id, "Track Updated", f"Your HVIA track is now {track.name}.", "enrollment")
    log(db, u, "INTERN_TRACK_CHANGED", "user", intern_id, {"track": track.name}, request); db.commit(); return {"ok": True}

@app.patch("/api/interns/{intern_id}/status/{status}")
def change_intern_status(intern_id: int, status: str, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin"))):
    intern = db.get(User, intern_id)
    allowed = {"onboarding", "active", "on_hold", "completed", "dropped", "terminated"}
    if not intern or intern.role != "intern" or status not in allowed: raise HTTPException(400)
    intern.status = status; intern.active = status not in {"terminated", "dropped"}
    log(db, u, "INTERN_STATUS_CHANGED", "user", intern_id, {"status": status}, request); db.commit(); return {"ok": True}

@app.delete("/api/interns/{intern_id}")
def delete_intern(intern_id: int, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin"))):
    intern = db.get(User, intern_id)
    if not intern or intern.role != "intern": raise HTTPException(404)
    intern.active = False; intern.status = "terminated"; intern.deleted_at = datetime.utcnow()
    for row in db.query(MentorAssignment).filter(MentorAssignment.intern_id==intern.id,MentorAssignment.active==True).all(): row.active=False
    for row in db.query(Enrollment).filter(Enrollment.user_id==intern.id,Enrollment.active==True).all(): row.active=False
    log(db, u, "INTERN_ARCHIVED", "user", intern_id, request=request); db.commit(); return {"ok": True}


# ---------- Tasks / work / submissions ----------
def ensure_task_assignments(db: Session, task: Task, actor: User, explicit_user_ids: list[int] | None = None):
    if explicit_user_ids:
        allowed_q = db.query(User).filter(User.role == "intern", User.active == True, User.deleted_at.is_(None), User.id.in_(explicit_user_ids))
        if actor.role == "admin":
            allowed_q = allowed_q.filter(User.id.in_(mentor_intern_ids(db, actor.id) or [-1]))
        users = allowed_q.all()
        # Never silently accept unauthorized/nonexistent IDs.
        if len({x.id for x in users}) != len(set(explicit_user_ids)):
            raise HTTPException(403, "One or more selected interns are outside your allowed scope")
    else:
        users = target_users_for_task(db, task, actor)
    for intern in users:
        if not db.query(TaskAssignment).filter(TaskAssignment.task_id == task.id, TaskAssignment.user_id == intern.id).first():
            db.add(TaskAssignment(task_id=task.id, user_id=intern.id))
            db.add(TaskProgress(task_id=task.id, user_id=intern.id, status="To Do"))
            notify(db, intern.id, "New Task Assigned", task.title, "task", "task", task.id)
    return len(users)

@app.get("/api/tasks")
def list_tasks(db: Session = Depends(get_db), u: User = Depends(current_user)):
    if u.role == "intern":
        ass = db.query(TaskAssignment).filter(TaskAssignment.user_id == u.id).all(); ids = [x.task_id for x in ass]
        rows = db.query(Task).filter(Task.id.in_(ids or [-1]), Task.archived_at.is_(None)).order_by(Task.id.desc()).all()
    elif u.role == "admin":
        rows = db.query(Task).filter(Task.created_by == u.id, Task.archived_at.is_(None)).order_by(Task.id.desc()).all()
    else:
        rows = db.query(Task).filter(Task.archived_at.is_(None)).order_by(Task.id.desc()).all()
    out = []
    for t in rows:
        progress = None; submission = None
        if u.role == "intern":
            p = db.query(TaskProgress).filter(TaskProgress.task_id == t.id, TaskProgress.user_id == u.id).first(); progress = p.status if p else "To Do"
            s = latest_submission(db, u.id, t.id); submission = serialize_submission(db, s) if s else None
        attachments = db.query(TaskAttachment).filter(TaskAttachment.task_id == t.id).all()
        assignment_rows = db.query(TaskAssignment).filter(TaskAssignment.task_id == t.id).all()
        assigned_users = [db.get(User, a.user_id) for a in assignment_rows]
        assigned_users = [x for x in assigned_users if x]
        track_name = db.get(Track, t.target_track_id).name if t.target_track_id and db.get(Track, t.target_track_id) else None
        out.append({"id": t.id, "title": t.title, "description": t.description, "instructions": t.instructions, "priority": t.priority, "deadline": t.deadline, "max_score": t.max_score, "target_type": t.target_type, "target_track_id": t.target_track_id, "target_track_name": track_name, "target_user_id": t.target_user_id, "created_by": t.created_by, "created_at": t.created_at, "assigned_count": len(assigned_users), "assigned_students": [{"id": x.id, "name": x.full_name, "email": x.email} for x in assigned_users], "progress": progress, "submission": submission, "attachments": [{"id": a.id, "name": a.original_name, "size": a.file_size, "download_url": f"/api/task-attachments/{a.id}/download"} for a in attachments]})
    return out

@app.post("/api/tasks")
def create_task(payload: TaskIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin", "admin"))):
    explicit_ids = list(dict.fromkeys(payload.target_user_ids or []))
    if payload.target_user_id and payload.target_user_id not in explicit_ids:
        explicit_ids.append(payload.target_user_id)

    # UI rule: selected students win; otherwise selected track means whole track; blank track means all allowed interns.
    if explicit_ids:
        target_type = "specific"
        target_track_id = payload.target_track_id
        target_user_id = explicit_ids[0] if len(explicit_ids) == 1 else None
    elif payload.target_track_id:
        target_type = "track"
        target_track_id = payload.target_track_id
        target_user_id = None
    else:
        target_type = "all"
        target_track_id = None
        target_user_id = None

    if target_track_id:
        track = db.get(Track, target_track_id)
        if not track or not track.active:
            raise HTTPException(400, "Invalid track")
        if explicit_ids:
            enrolled_ids = {e.user_id for e in db.query(Enrollment).filter(Enrollment.track_id == target_track_id, Enrollment.active == True, Enrollment.user_id.in_(explicit_ids)).all()}
            if enrolled_ids != set(explicit_ids):
                raise HTTPException(400, "Selected interns must belong to the selected track")

    data = payload.model_dump(exclude={"target_user_ids"})
    data.update({"target_type": target_type, "target_track_id": target_track_id, "target_user_id": target_user_id})
    task = Task(**data, created_by=u.id)
    db.add(task); db.flush()
    count = ensure_task_assignments(db, task, u, explicit_ids if explicit_ids else None)
    if count == 0:
        db.rollback()
        raise HTTPException(400, "No interns matched this assignment")
    log(db, u, "TASK_CREATED", "task", task.id, {
        "title": task.title, "assigned": count, "target_type": target_type,
        "track_id": target_track_id, "selected_user_ids": explicit_ids
    }, request)
    db.commit()
    return {"id": task.id, "assigned": count, "target_type": target_type}

@app.patch("/api/tasks/{task_id}")
def update_task(task_id:int,payload:TaskIn,request:Request,db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin","admin"))):
    task=db.get(Task,task_id)
    if not task or task.archived_at is not None: raise HTTPException(404,"Task not found")
    if u.role=="admin" and task.created_by!=u.id: raise HTTPException(403,"Admins can edit only tasks they created")
    explicit_ids=list(dict.fromkeys(payload.target_user_ids or []))
    if payload.target_user_id and payload.target_user_id not in explicit_ids: explicit_ids.append(payload.target_user_id)
    if explicit_ids:
        target_type="specific"; target_track_id=payload.target_track_id; target_user_id=explicit_ids[0] if len(explicit_ids)==1 else None
    elif payload.target_track_id:
        target_type="track"; target_track_id=payload.target_track_id; target_user_id=None
    else:
        target_type="all"; target_track_id=None; target_user_id=None
    if target_track_id:
        track=db.get(Track,target_track_id)
        if not track or not track.active: raise HTTPException(400,"Invalid track")
        if explicit_ids:
            valid={e.user_id for e in db.query(Enrollment).filter(Enrollment.track_id==target_track_id,Enrollment.active==True,Enrollment.user_id.in_(explicit_ids)).all()}
            if valid!=set(explicit_ids): raise HTTPException(400,"Selected interns must belong to the selected track")
    for field in ["title","description","instructions","priority","deadline","max_score","grace_minutes","allow_resubmission"]: setattr(task,field,getattr(payload,field))
    task.target_type=target_type;task.target_track_id=target_track_id;task.target_user_id=target_user_id
    db.flush()
    desired_users = []
    if explicit_ids:
        q=db.query(User).filter(User.role=="intern",User.active==True,User.deleted_at.is_(None),User.id.in_(explicit_ids))
        if u.role=="admin": q=q.filter(User.id.in_(mentor_intern_ids(db,u.id) or [-1]))
        desired_users=q.all()
        if {x.id for x in desired_users} != set(explicit_ids):
            raise HTTPException(403, "One or more selected interns are outside your allowed scope")
    else: desired_users=target_users_for_task(db,task,u)
    desired={x.id for x in desired_users}; existing_rows=db.query(TaskAssignment).filter(TaskAssignment.task_id==task.id).all(); existing={x.user_id for x in existing_rows}
    added=desired-existing; removed=existing-desired; retained=[]
    for iid in added:
        db.add(TaskAssignment(task_id=task.id,user_id=iid));db.add(TaskProgress(task_id=task.id,user_id=iid,status="To Do"));notify(db,iid,"Task Assigned",task.title,"task","task",task.id)
    for iid in removed:
        if db.query(Submission).filter(Submission.task_id==task.id,Submission.user_id==iid).first(): retained.append(iid);continue
        db.query(TaskAssignment).filter(TaskAssignment.task_id==task.id,TaskAssignment.user_id==iid).delete(synchronize_session=False);db.query(TaskProgress).filter(TaskProgress.task_id==task.id,TaskProgress.user_id==iid).delete(synchronize_session=False)
    if not (desired or retained):
        db.rollback()
        raise HTTPException(400, "No interns matched this assignment")
    notify_many(db,list(desired|set(retained)),"Task Updated",task.title,"task","task",task.id)
    log(db,u,"TASK_UPDATED","task",task.id,{"added":list(added),"removed":list(removed-set(retained)),"retained_with_submission":retained},request);db.commit();return {"ok":True,"assigned":len(desired|set(retained)),"retained_with_submission":retained}

@app.delete("/api/tasks/{task_id}")
def delete_task(task_id:int,request:Request,db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin","admin"))):
    task=db.get(Task,task_id)
    if not task or task.archived_at is not None: raise HTTPException(404,"Task not found")
    if u.role=="admin" and task.created_by!=u.id: raise HTTPException(403,"Admins can delete only tasks they created")
    task.archived_at=datetime.utcnow();task.status="archived"
    ass=[x.user_id for x in db.query(TaskAssignment).filter(TaskAssignment.task_id==task.id).all()]
    notify_many(db,ass,"Task Archived",task.title,"task","task",task.id);log(db,u,"TASK_ARCHIVED","task",task.id,{"title":task.title},request);db.commit();return {"ok":True}

@app.delete("/api/task-attachments/{attachment_id}")
def delete_task_attachment(attachment_id:int,request:Request,db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin","admin"))):
    a=db.get(TaskAttachment,attachment_id)
    if not a: raise HTTPException(404)
    t=db.get(Task,a.task_id)
    if u.role=="admin" and t.created_by!=u.id: raise HTTPException(403)
    path=UPLOADS/"task_attachments"/a.stored_name
    if path.exists(): path.unlink(missing_ok=True)
    db.delete(a);log(db,u,"TASK_ATTACHMENT_DELETED","task",t.id,{"name":a.original_name},request);db.commit();return {"ok":True}

@app.post("/api/tasks/{task_id}/attachments")
async def add_task_attachments(task_id: int, uploads: list[UploadFile] = File(...), request: Request = None, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin", "admin"))):
    t = db.get(Task, task_id)
    if not t or (u.role == "admin" and t.created_by != u.id): raise HTTPException(404)
    added = []
    for upload in uploads:
        original, stored, size = save_upload(upload, UPLOADS/"task_attachments")
        a = TaskAttachment(task_id=t.id, original_name=original, stored_name=stored, content_type=upload.content_type, file_size=size, uploaded_by=u.id)
        db.add(a); db.flush(); added.append({"id": a.id, "name": original, "size": size})
    log(db, u, "TASK_FILES_UPLOADED", "task", task_id, {"files": [x["name"] for x in added]}, request); db.commit(); return added

@app.get("/api/task-attachments/{attachment_id}/download")
def task_attachment_download(attachment_id: int, db: Session = Depends(get_db), u: User = Depends(current_user)):
    a = db.get(TaskAttachment, attachment_id)
    if not a: raise HTTPException(404)
    t = db.get(Task, a.task_id)
    if u.role == "intern" and not task_for_user(db, t.id, u.id): raise HTTPException(403)
    if u.role == "admin" and t.created_by != u.id and u.id not in [t.created_by]: raise HTTPException(403)
    path = UPLOADS/"task_attachments"/a.stored_name
    if not path.exists(): raise HTTPException(404)
    return FileResponse(path, filename=a.original_name, media_type=a.content_type or "application/octet-stream")

@app.patch("/api/tasks/{task_id}/status")
def update_task_status(task_id: int, payload: dict, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("intern"))):
    p = db.query(TaskProgress).filter(TaskProgress.task_id == task_id, TaskProgress.user_id == u.id).first()
    if not p: raise HTTPException(404)
    status = payload.get("status")
    if status not in {"To Do", "In Progress"}: raise HTTPException(400, "Intern can only move between To Do and In Progress. Submission changes status automatically.")
    p.status = status; log(db, u, "TASK_STATUS_CHANGED", "task", task_id, {"status": status}, request); db.commit(); return {"ok": True}

@app.post("/api/work/start/{task_id}")
def work_start(task_id: int, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("intern"))):
    if not task_for_user(db, task_id, u.id): raise HTTPException(403)
    open_session = db.query(WorkSession).filter(WorkSession.user_id == u.id, WorkSession.ended_at.is_(None)).first()
    if open_session: raise HTTPException(400, "You already have an active work timer")
    ws = WorkSession(user_id=u.id, task_id=task_id, started_at=datetime.utcnow())
    db.add(ws); log(db, u, "WORK_TIMER_STARTED", "task", task_id, request=request); db.commit(); db.refresh(ws); return {"id": ws.id, "started_at": ws.started_at}

@app.post("/api/work/stop")
def work_stop(payload: dict, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("intern"))):
    ws = db.query(WorkSession).filter(WorkSession.user_id == u.id, WorkSession.ended_at.is_(None)).order_by(WorkSession.id.desc()).first()
    if not ws: raise HTTPException(404, "No active timer")
    ws.ended_at = datetime.utcnow(); ws.minutes = max(1, int((ws.ended_at-ws.started_at).total_seconds()/60)); ws.note = payload.get("note", "")
    log(db, u, "WORK_TIMER_STOPPED", "task", ws.task_id, {"minutes": ws.minutes}, request); db.commit(); return {"minutes": ws.minutes, "total_hours": hours_for_user(db, u.id)}

@app.get("/api/work")
def work_summary(db: Session = Depends(get_db), u: User = Depends(current_user)):
    target = u.id
    active = db.query(WorkSession).filter(WorkSession.user_id == target, WorkSession.ended_at.is_(None)).first()
    return {"hours": hours_for_user(db, target), "active": {"id": active.id, "task_id": active.task_id, "started_at": active.started_at} if active else None}

@app.post("/api/tasks/{task_id}/submit")
async def submit_task(task_id: int, notes: str = Form(""), github_url: str = Form(""), drive_url: str = Form(""), uploads: list[UploadFile] = File(default=[]), request: Request = None, db: Session = Depends(get_db), u: User = Depends(require_roles("intern"))):
    if not task_for_user(db, task_id, u.id): raise HTTPException(403)
    t = db.get(Task, task_id); old = latest_submission(db, u.id, task_id); version = (old.version + 1) if old else 1
    if old and not t.allow_resubmission and old.status != "Changes Requested": raise HTTPException(400, "Resubmission is disabled")
    late = bool(t.deadline and datetime.utcnow() > t.deadline + timedelta(minutes=t.grace_minutes or 0))
    s = Submission(task_id=task_id, user_id=u.id, version=version, notes=notes, github_url=github_url or None, drive_url=drive_url or None, status="Submitted", is_late=late)
    db.add(s); db.flush()
    for upload in uploads:
        if not upload.filename: continue
        original, stored, size = save_upload(upload, UPLOADS/"submissions")
        db.add(SubmissionFile(submission_id=s.id, original_name=original, stored_name=stored, content_type=upload.content_type, file_size=size))
    p = db.query(TaskProgress).filter(TaskProgress.task_id == task_id, TaskProgress.user_id == u.id).first()
    if p: p.status = "Submitted"
    creator = db.get(User, t.created_by); notify(db, t.created_by, "New Submission", f"{u.full_name} submitted {t.title} (v{version}).", "submission", "submission", s.id)
    log(db, u, "TASK_SUBMITTED", "submission", s.id, {"task": t.title, "version": version, "late": late}, request); db.commit(); return serialize_submission(db, s)

@app.get("/api/submissions")
def submissions(db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin", "admin"))):
    q = db.query(Submission)
    if u.role == "admin":
        scope = mentor_intern_ids(db, u.id); q = q.filter(Submission.user_id.in_(scope or [-1]))
    return [serialize_submission(db, s) for s in q.order_by(Submission.id.desc()).all()]

@app.get("/api/submission-files/{file_id}/download")
def submission_file_download(file_id: int, db: Session = Depends(get_db), u: User = Depends(current_user)):
    f = db.get(SubmissionFile, file_id)
    if not f: raise HTTPException(404)
    s = db.get(Submission, f.submission_id)
    if u.role == "intern" and s.user_id != u.id: raise HTTPException(403)
    if u.role == "admin" and s.user_id not in mentor_intern_ids(db, u.id): raise HTTPException(403)
    path = UPLOADS/"submissions"/f.stored_name
    if not path.exists(): raise HTTPException(404)
    return FileResponse(path, filename=f.original_name, media_type=f.content_type or "application/octet-stream")

@app.post("/api/submissions/{submission_id}/review")
def review_submission(submission_id: int, payload: ReviewIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin", "admin"))):
    s = db.get(Submission, submission_id)
    if not s: raise HTTPException(404)
    if u.role == "admin" and s.user_id not in mentor_intern_ids(db, u.id): raise HTTPException(403)
    t = db.get(Task, s.task_id)
    if payload.score is not None and (payload.score < 0 or payload.score > t.max_score): raise HTTPException(400, "Invalid score")
    s.score = payload.score; s.feedback = payload.feedback; s.status = payload.status; s.reviewed_by = u.id; s.reviewed_at = datetime.utcnow()
    p = db.query(TaskProgress).filter(TaskProgress.task_id == s.task_id, TaskProgress.user_id == s.user_id).first()
    if p: p.status = payload.status
    notify(db, s.user_id, "Task Review Updated", f"{t.title}: {payload.status}" + (f" · Score {payload.score}/{t.max_score}" if payload.score is not None else ""), "review", "submission", s.id)
    log(db, u, "SUBMISSION_REVIEWED", "submission", s.id, {"student_id": s.user_id, "status": payload.status, "score": payload.score}, request); db.commit(); return {"ok": True}


# ---------- Meetings / attendance ----------
def meetings_for_user(db: Session, u: User):
    if u.role == "super_admin": return db.query(Meeting).order_by(Meeting.starts_at.desc()).all()
    ids = [x.meeting_id for x in db.query(MeetingAttendee).filter(MeetingAttendee.user_id == u.id).all()]
    if u.role == "admin": ids += [x.id for x in db.query(Meeting).filter(Meeting.created_by == u.id).all()]
    return db.query(Meeting).filter(Meeting.id.in_(list(set(ids)) or [-1])).order_by(Meeting.starts_at.desc()).all()

@app.get("/api/meetings")
def meetings(db: Session = Depends(get_db), u: User = Depends(current_user)):
    out=[]
    for m in meetings_for_user(db,u):
        attendees = db.query(MeetingAttendee).filter(MeetingAttendee.meeting_id == m.id).all()
        out.append({"id":m.id,"title":m.title,"description":m.description,"starts_at":m.starts_at,"duration_minutes":m.duration_minutes,"meeting_link":m.meeting_link,"target_type":m.target_type,"target_track_id":m.target_track_id,"target_user_id":m.target_user_id,"created_by":m.created_by,"attendee_count":len(attendees),"present_count":sum(1 for a in attendees if a.attendance_status in ("present","late"))})
    return out

@app.post("/api/meetings")
def create_meeting(payload: MeetingIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_roles("super_admin", "admin"))):
    tt=payload.target_type
    if tt=="track" and not payload.target_track_id: raise HTTPException(400,"Track is required for One Track meetings")
    if tt in {"track","specific"} and payload.target_track_id:
        chosen_track=db.get(Track,payload.target_track_id)
        if not chosen_track or not chosen_track.active: raise HTTPException(400,"Invalid or inactive track")
    if tt=="specific" and (not payload.target_track_id or not payload.target_user_id): raise HTTPException(400,"Choose a track and an intern for Specific Intern")
    if tt in {"admins","specific_admin","all"} and u.role!="super_admin": raise HTTPException(403,"Only Super Admin can schedule mentor/admin meetings")
    if tt=="specific":
        person=db.get(User,payload.target_user_id)
        if not person or person.role!="intern": raise HTTPException(400,"Invalid intern")
        e=active_enrollment(db,person.id)
        if not e or e.track_id!=payload.target_track_id: raise HTTPException(400,"Selected intern does not belong to the selected track")
        if u.role=="admin" and person.id not in mentor_intern_ids(db,u.id): raise HTTPException(403,"Intern outside your mentor scope")
    if tt=="specific_admin":
        person=db.get(User,payload.target_user_id)
        if not person or person.role!="admin" or not person.active: raise HTTPException(400,"Invalid mentor/admin")
    m=Meeting(**payload.model_dump(),created_by=u.id);db.add(m);db.flush()
    candidates=[]
    if u.role=="super_admin":
        if tt=="admins": candidates=[x.id for x in db.query(User).filter(User.role=="admin",User.active==True,User.deleted_at.is_(None)).all()]
        elif tt=="specific_admin": candidates=[payload.target_user_id]
        elif tt=="all": candidates=[x.id for x in db.query(User).filter(User.active==True,User.deleted_at.is_(None),User.role.in_(["intern","admin"])).all()]
        elif tt=="interns": candidates=[x.id for x in db.query(User).filter(User.role=="intern",User.active==True,User.deleted_at.is_(None)).all()]
        elif tt=="track": candidates=[e.user_id for e in db.query(Enrollment).filter(Enrollment.track_id==payload.target_track_id,Enrollment.active==True).all()]
        elif tt=="specific": candidates=[payload.target_user_id]
    else:
        scope=set(mentor_intern_ids(db,u.id))
        if tt=="interns": candidates=list(scope)
        elif tt=="track": candidates=[e.user_id for e in db.query(Enrollment).filter(Enrollment.track_id==payload.target_track_id,Enrollment.active==True).all() if e.user_id in scope]
        elif tt=="specific" and payload.target_user_id in scope: candidates=[payload.target_user_id]
    candidates=list(dict.fromkeys([x for x in candidates if x]))
    if not candidates: db.rollback();raise HTTPException(400,"No attendees matched this meeting audience")
    for uid in candidates: db.add(MeetingAttendee(meeting_id=m.id,user_id=uid,attendance_status="invited"))
    db.flush();notify_many(db,candidates,"New Meeting",f"{m.title} · {m.starts_at.strftime('%d %b %Y %H:%M')}","meeting","meeting",m.id);log(db,u,"MEETING_CREATED","meeting",m.id,{"target":tt,"attendees":len(candidates)},request);db.commit();return {"id":m.id,"attendees":len(candidates)}

@app.patch("/api/meetings/{meeting_id}")
def update_meeting(meeting_id:int,payload:MeetingIn,request:Request,db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin","admin"))):
    m=db.get(Meeting,meeting_id)
    if not m: raise HTTPException(404,"Meeting not found")
    if u.role=="admin" and m.created_by!=u.id: raise HTTPException(403,"Admins can edit only meetings they created")
    # Reuse creation validation/recipient logic by validating then rebuilding attendees locally.
    tt=payload.target_type
    if tt=="track" and not payload.target_track_id: raise HTTPException(400,"Track is required for One Track meetings")
    if tt in {"track","specific"} and payload.target_track_id:
        chosen_track=db.get(Track,payload.target_track_id)
        if not chosen_track or not chosen_track.active: raise HTTPException(400,"Invalid or inactive track")
    if tt=="specific" and (not payload.target_track_id or not payload.target_user_id): raise HTTPException(400,"Choose a track and an intern for Specific Intern")
    if tt in {"admins","specific_admin","all"} and u.role!="super_admin": raise HTTPException(403)
    if tt=="specific":
        e=active_enrollment(db,payload.target_user_id)
        if not e or e.track_id!=payload.target_track_id: raise HTTPException(400,"Selected intern does not belong to the selected track")
        if u.role=="admin" and payload.target_user_id not in mentor_intern_ids(db,u.id): raise HTTPException(403)
    if tt=="specific_admin":
        p=db.get(User,payload.target_user_id)
        if not p or p.role!="admin" or not p.active: raise HTTPException(400,"Invalid mentor/admin")
    for f,v in payload.model_dump().items(): setattr(m,f,v)
    db.query(MeetingAttendee).filter(MeetingAttendee.meeting_id==m.id).delete(synchronize_session=False)
    if u.role=="super_admin":
        if tt=="admins": cand=[x.id for x in db.query(User).filter(User.role=="admin",User.active==True,User.deleted_at.is_(None)).all()]
        elif tt=="specific_admin": cand=[payload.target_user_id]
        elif tt=="all": cand=[x.id for x in db.query(User).filter(User.active==True,User.deleted_at.is_(None),User.role.in_(["intern","admin"])).all()]
        elif tt=="interns": cand=[x.id for x in db.query(User).filter(User.role=="intern",User.active==True,User.deleted_at.is_(None)).all()]
        elif tt=="track": cand=[e.user_id for e in db.query(Enrollment).filter(Enrollment.track_id==payload.target_track_id,Enrollment.active==True).all()]
        else: cand=[payload.target_user_id]
    else:
        scope=set(mentor_intern_ids(db,u.id));cand=list(scope) if tt=="interns" else ([e.user_id for e in db.query(Enrollment).filter(Enrollment.track_id==payload.target_track_id,Enrollment.active==True).all() if e.user_id in scope] if tt=="track" else [payload.target_user_id])
    cand=list(dict.fromkeys([x for x in cand if x]))
    if not cand: db.rollback();raise HTTPException(400,"No attendees matched this meeting audience")
    for uid in cand: db.add(MeetingAttendee(meeting_id=m.id,user_id=uid,attendance_status="invited"))
    notify_many(db,cand,"Meeting Updated",f"{m.title} · {m.starts_at.strftime('%d %b %Y %H:%M')}","meeting","meeting",m.id);log(db,u,"MEETING_UPDATED","meeting",m.id,{"target":tt,"attendees":len(cand)},request);db.commit();return {"ok":True,"attendees":len(cand)}

@app.delete("/api/meetings/{meeting_id}")
def delete_meeting(meeting_id:int,request:Request,db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin","admin"))):
    m=db.get(Meeting,meeting_id)
    if not m: raise HTTPException(404,"Meeting not found")
    if u.role=="admin" and m.created_by!=u.id: raise HTTPException(403)
    ids=[x.user_id for x in db.query(MeetingAttendee).filter(MeetingAttendee.meeting_id==m.id).all()]
    db.query(MeetingAttendee).filter(MeetingAttendee.meeting_id==m.id).delete(synchronize_session=False);log(db,u,"MEETING_DELETED","meeting",m.id,{"title":m.title},request);db.delete(m);notify_many(db,ids,"Meeting Cancelled",m.title,"meeting");db.commit();return {"ok":True}

@app.get("/api/meetings/{meeting_id}/attendance")
def meeting_attendance(meeting_id:int, db:Session=Depends(get_db), u:User=Depends(require_roles("super_admin","admin"))):
    rows=db.query(MeetingAttendee).filter(MeetingAttendee.meeting_id==meeting_id).all(); out=[]
    for x in rows:
        person=db.get(User,x.user_id)
        if u.role=="admin" and person.role=="intern" and person.id not in mentor_intern_ids(db,u.id): continue
        out.append({"id":x.id,"user_id":x.user_id,"name":person.full_name,"role":person.role,"status":x.attendance_status,"check_in_at":x.check_in_at,"note":x.note})
    return out

@app.patch("/api/meetings/attendance/{attendee_id}")
def update_attendance(attendee_id:int,payload:AttendanceIn,request:Request,db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin","admin"))):
    a=db.get(MeetingAttendee,attendee_id)
    if not a: raise HTTPException(404)
    person=db.get(User,a.user_id)
    if u.role=="admin" and person.role=="intern" and person.id not in mentor_intern_ids(db,u.id): raise HTTPException(403)
    a.attendance_status=payload.status; a.note=payload.note; a.check_in_at=datetime.utcnow() if payload.status in ("present","late") else None
    log(db,u,"ATTENDANCE_UPDATED","meeting_attendee",a.id,{"user_id":a.user_id,"status":payload.status},request); db.commit(); return {"ok":True}


# ---------- Announcements / resources ----------
@app.get("/api/announcements")
def announcements(db:Session=Depends(get_db),u:User=Depends(current_user)):
    now=datetime.utcnow(); rows=db.query(Announcement).filter(or_(Announcement.publish_at.is_(None),Announcement.publish_at<=now)).order_by(Announcement.pinned.desc(),Announcement.id.desc()).all(); out=[]
    e=active_enrollment(db,u.id)
    for a in rows:
        visible=a.target_type=="all" or (a.target_type=="role" and a.target_role==u.role) or (a.target_type=="track" and e and e.track_id==a.target_track_id)
        if u.role=="super_admin": visible=True
        if visible: out.append({"id":a.id,"title":a.title,"content":a.content,"pinned":a.pinned,"created_at":a.created_at,"publish_at":a.publish_at,"target_type":a.target_type})
    return out

@app.post("/api/announcements")
def create_announcement(payload:AnnouncementIn,request:Request,db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin","admin"))):
    a=Announcement(**payload.model_dump(),created_by=u.id);db.add(a);db.flush();log(db,u,"ANNOUNCEMENT_CREATED","announcement",a.id,{"title":a.title},request);db.commit();return {"id":a.id}

@app.get("/api/resources")
def resources(db:Session=Depends(get_db),u:User=Depends(current_user)):
    q=db.query(LearningResource)
    if u.role=="intern":
        e=active_enrollment(db,u.id); q=q.filter(LearningResource.track_id==e.track_id) if e and e.track_id else q.filter(LearningResource.id==-1)
    rows=q.order_by(LearningResource.track_id,LearningResource.order_index).all();out=[]
    for r in rows:
        p=db.query(ResourceProgress).filter(ResourceProgress.resource_id==r.id,ResourceProgress.user_id==u.id).first() if u.role=="intern" else None
        out.append({"id":r.id,"track_id":r.track_id,"title":r.title,"resource_type":r.resource_type,"url":r.url,"description":r.description,"module_name":r.module_name,"order_index":r.order_index,"progress":p.progress_percent if p else 0,"completed":p.completed if p else False})
    return out

@app.post("/api/resources/upload")
async def upload_resource(track_id:int=Form(...),title:str=Form(...),resource_type:str=Form("file"),module_name:str=Form(""),description:str=Form(""),upload:UploadFile=File(...),request:Request=None,db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin","admin"))):
    original,stored,size=save_upload(upload,UPLOADS/"resources");r=LearningResource(track_id=track_id,title=title,resource_type=resource_type,url=f"/api/resources/{stored}/download",description=description,module_name=module_name,created_by=u.id);db.add(r);db.flush();log(db,u,"RESOURCE_UPLOADED","resource",r.id,{"file":original,"size":size},request);db.commit();return {"id":r.id}

@app.get("/api/resources/{stored}/download")
def resource_download(stored:str,db:Session=Depends(get_db),u:User=Depends(current_user)):
    path=UPLOADS/"resources"/Path(stored).name
    if not path.exists():raise HTTPException(404)
    return FileResponse(path)

@app.post("/api/resources/{resource_id}/progress")
def update_resource_progress(resource_id:int,payload:ResourceProgressIn,db:Session=Depends(get_db),u:User=Depends(require_roles("intern"))):
    r=db.get(LearningResource,resource_id)
    if not r:raise HTTPException(404)
    p=db.query(ResourceProgress).filter(ResourceProgress.resource_id==resource_id,ResourceProgress.user_id==u.id).first()
    if not p:p=ResourceProgress(resource_id=resource_id,user_id=u.id);db.add(p)
    p.watched_seconds=max(p.watched_seconds,payload.watched_seconds);p.progress_percent=max(p.progress_percent,min(100,payload.progress_percent));p.completed=p.completed or payload.completed or p.progress_percent>=95;p.last_opened_at=datetime.utcnow();db.commit();return {"ok":True}


# ---------- Notifications / complaints / warnings ----------
@app.get("/api/notifications")
def notifications(db:Session=Depends(get_db),u:User=Depends(current_user)):
    rows=db.query(Notification).filter(Notification.user_id==u.id).order_by(Notification.id.desc()).limit(100).all();return [{"id":x.id,"title":x.title,"message":x.message,"type":x.notification_type,"is_read":x.is_read,"created_at":x.created_at} for x in rows]

@app.patch("/api/notifications/{notification_id}/read")
def notification_read(notification_id:int,db:Session=Depends(get_db),u:User=Depends(current_user)):
    n=db.get(Notification,notification_id)
    if not n or n.user_id!=u.id:raise HTTPException(404)
    n.is_read=True;db.commit();return {"ok":True}

@app.post("/api/complaints")
def create_complaint(payload:ComplaintIn,request:Request,db:Session=Depends(get_db),u:User=Depends(require_roles("intern"))):
    c=Complaint(user_id=u.id,**payload.model_dump());db.add(c);db.flush();super_ids=[x.id for x in db.query(User).filter(User.role=="super_admin",User.active==True).all()];notify_many(db,super_ids,"New Intern Complaint",f"{u.full_name}: {payload.subject}","complaint","complaint",c.id);log(db,u,"COMPLAINT_CREATED","complaint",c.id,{"subject":payload.subject},request);db.commit();return {"id":c.id}

@app.get("/api/complaints")
def complaints(db:Session=Depends(get_db),u:User=Depends(current_user)):
    q=db.query(Complaint)
    if u.role=="intern":q=q.filter(Complaint.user_id==u.id)
    elif u.role!="super_admin":raise HTTPException(403)
    out=[]
    for c in q.order_by(Complaint.id.desc()).all():
        student=db.get(User,c.user_id);out.append({"id":c.id,"student":student.full_name if student else "—","subject":c.subject,"message":c.message,"category":c.category,"priority":c.priority,"status":c.status,"reply":c.super_admin_reply,"created_at":c.created_at})
    return out

@app.post("/api/complaints/{complaint_id}/reply")
def complaint_reply(complaint_id:int,payload:ComplaintReplyIn,request:Request,db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin"))):
    c=db.get(Complaint,complaint_id)
    if not c:raise HTTPException(404)
    c.status=payload.status;c.super_admin_reply=payload.reply;c.resolved_at=datetime.utcnow() if payload.status.lower() in ("resolved","closed") else None;notify(db,c.user_id,"Complaint Updated",payload.reply or f"Status: {payload.status}","complaint","complaint",c.id);log(db,u,"COMPLAINT_UPDATED","complaint",c.id,{"status":payload.status},request);db.commit();return {"ok":True}

@app.post("/api/warnings")
def create_warning(payload:WarningIn,request:Request,db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin","admin"))):
    if u.role=="admin" and payload.user_id not in mentor_intern_ids(db,u.id):raise HTTPException(403)
    w=Warning(user_id=payload.user_id,created_by=u.id,warning_type=payload.warning_type,severity=payload.severity,note=payload.note);db.add(w);db.flush();notify(db,payload.user_id,"Performance Warning",payload.note,"warning","warning",w.id);log(db,u,"WARNING_CREATED","warning",w.id,{"user_id":payload.user_id,"severity":payload.severity},request);db.commit();return {"id":w.id}


# ---------- Monitoring ----------
@app.get("/api/monitoring")
def monitoring(db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin"))):
    now=datetime.utcnow();today=now.replace(hour=0,minute=0,second=0,microsecond=0);week=now-timedelta(days=7)
    interns=db.query(User).filter(User.role=="intern",User.active==True,User.deleted_at.is_(None)).count();admins=db.query(User).filter(User.role=="admin",User.active==True,User.deleted_at.is_(None)).count();logins=db.query(LoginEvent).filter(LoginEvent.success==True,LoginEvent.created_at>=today).count();online=db.query(User).filter(User.last_seen_at>=now-timedelta(minutes=10),User.active==True).count();subs=db.query(Submission).filter(Submission.submitted_at>=today).count();late=db.query(Submission).filter(Submission.is_late==True,Submission.submitted_at>=week).count();pending=db.query(Submission).filter(Submission.status.in_(["Submitted","Under Review"])).count();meet=db.query(Meeting).filter(Meeting.starts_at>=today,Meeting.starts_at<today+timedelta(days=1)).count();complaints_open=db.query(Complaint).filter(Complaint.status.in_(["Open","In Progress"])).count()
    at_risk=[]
    for person in db.query(User).filter(User.role=="intern",User.active==True).all():
        stats=task_stats_for_intern(db,person.id);att=attendance_for_user(db,person.id);reasons=[]
        if stats["late"]>=3:reasons.append("3+ late submissions")
        if att["total"]>=3 and att["percent"]<70:reasons.append("attendance below 70%")
        if person.last_login_at and person.last_login_at<now-timedelta(days=7):reasons.append("inactive for 7+ days")
        if reasons:at_risk.append({"id":person.id,"name":person.full_name,"reasons":reasons})
    events=[]
    for x in db.query(ActivityLog).order_by(ActivityLog.id.desc()).limit(30).all():
        actor=db.get(User,x.actor_user_id) if x.actor_user_id else None;events.append({"id":x.id,"actor":actor.full_name if actor else "System","action":x.action,"entity_type":x.entity_type,"details":x.details,"created_at":x.created_at})
    tracks=[]
    for t in db.query(Track).filter(Track.active==True).all():tracks.append({"id":t.id,"name":t.name,"count":db.query(Enrollment).filter(Enrollment.track_id==t.id,Enrollment.active==True).count()})
    return {"kpis":{"interns":interns,"admins":admins,"online_users":online,"logins_today":logins,"submissions_today":subs,"late_week":late,"pending_reviews":pending,"meetings_today":meet,"open_complaints":complaints_open,"at_risk":len(at_risk)},"tracks":tracks,"at_risk":at_risk,"events":events}

@app.get("/api/global-search")
def global_search(q:str=Query(min_length=1),db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin"))):
    term=f"%{q}%";users=db.query(User).filter(or_(User.full_name.ilike(term),User.email.ilike(term),User.internal_id.ilike(term))).limit(12).all();tasks=db.query(Task).filter(Task.title.ilike(term),Task.archived_at.is_(None)).limit(12).all();return {"users":[serialize_user(db,x) for x in users],"tasks":[{"id":x.id,"title":x.title} for x in tasks]}


# ---------- User preferences ----------
@app.get("/api/preferences")
def get_preferences(db: Session = Depends(get_db), u: User = Depends(current_user)):
    pref = db.query(UserPreference).filter(UserPreference.user_id == u.id).first()
    return {"theme": pref.theme if pref else "dark"}

@app.patch("/api/preferences/theme")
def set_theme(payload: SettingIn, db: Session = Depends(get_db), u: User = Depends(current_user)):
    theme = payload.value.lower().strip()
    if theme not in {"dark", "light"}:
        raise HTTPException(400, "Theme must be dark or light")
    pref = db.query(UserPreference).filter(UserPreference.user_id == u.id).first()
    if not pref:
        pref = UserPreference(user_id=u.id, theme=theme)
        db.add(pref)
    else:
        pref.theme = theme
    db.commit()
    return {"theme": theme}


# ---------- Dashboard / settings / activity ----------
@app.get("/api/dashboard")
def dashboard(db:Session=Depends(get_db),u:User=Depends(current_user)):
    if u.role=="super_admin":
        m=monitoring(db,u);return m
    if u.role=="admin":
        scope=mentor_intern_ids(db,u.id);pending=db.query(Submission).filter(Submission.user_id.in_(scope or [-1]),Submission.status.in_(["Submitted","Under Review"])).count();reviews=db.query(Submission).filter(Submission.reviewed_by==u.id).count();return {"kpis":{"assigned_interns":len(scope),"pending_reviews":pending,"reviews_done":reviews,"tasks_created":db.query(Task).filter(Task.created_by==u.id).count(),"meetings":db.query(Meeting).filter(Meeting.created_by==u.id).count()},"recent_submissions":[serialize_submission(db,s) for s in db.query(Submission).filter(Submission.user_id.in_(scope or [-1])).order_by(Submission.id.desc()).limit(8).all()]}
    stats=task_stats_for_intern(db,u.id);att=attendance_for_user(db,u.id);notifications_count=db.query(Notification).filter(Notification.user_id==u.id,Notification.is_read==False).count();return {"kpis":{**stats,"hours":hours_for_user(db,u.id),"attendance":att["percent"],"notifications":notifications_count},"track":serialize_user(db,u)["track"]}

@app.get("/api/activity")
def activity(db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin"))):
    rows=db.query(ActivityLog).order_by(ActivityLog.id.desc()).limit(200).all();return [{"id":x.id,"actor_user_id":x.actor_user_id,"action":x.action,"entity_type":x.entity_type,"entity_id":x.entity_id,"details":x.details,"ip_address":x.ip_address,"created_at":x.created_at} for x in rows]

@app.get("/api/settings")
def settings(db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin"))):
    return {x.key:x.value for x in db.query(SystemSetting).all()}

@app.patch("/api/settings/{key}")
def update_setting(key:str,payload:SettingIn,request:Request,db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin"))):
    s=db.query(SystemSetting).filter(SystemSetting.key==key).first()
    if not s:s=SystemSetting(key=key,value=payload.value,updated_by=u.id);db.add(s)
    else:s.value=payload.value;s.updated_by=u.id
    log(db,u,"SETTING_UPDATED","setting",None,{"key":key,"value":payload.value},request);db.commit();return {"ok":True}

@app.get("/api/contact-requests")
def contact_requests(db:Session=Depends(get_db),u:User=Depends(require_roles("super_admin"))):
    return [{"id":x.id,"name":x.name,"phone":x.phone,"email":x.email,"company":x.company,"interest":x.interest,"message":x.message,"status":x.status,"created_at":x.created_at} for x in db.query(ContactRequest).order_by(ContactRequest.id.desc()).all()]


# demo placeholders
@app.get("/demo-video",response_class=HTMLResponse)
def demo_video():return "<html><body style='background:#061525;color:white;font-family:Arial;display:grid;place-items:center;height:100vh'><div><h1>HVIA Protected Video Placeholder</h1><p>Connect Cloudflare Stream / Bunny Stream for production video protection.</p></div></body></html>"
@app.get("/demo-book",response_class=HTMLResponse)
def demo_book():return "<html><body style='font-family:Arial;padding:40px'><h1>HVIA Learning Resource</h1><p>Replace this placeholder with your uploaded book or PDF.</p></body></html>"
